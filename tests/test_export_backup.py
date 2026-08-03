from pathlib import Path

from openpyxl import load_workbook

from src.services.backup_service import create_database_backup
from src.services.export_service import export_vouchers_to_excel


def test_export_vouchers_to_excel_creates_workbook(tmp_path: Path):
    output_path = tmp_path / "vales.xlsx"
    vouchers = [
        {
            "serial_number": "VAL-202401-0001",
            "issue_date": "2024-01-10",
            "employee_name": "Ana",
            "area": "Mantenimiento",
            "vehicle_id": 1,
            "fuel_type": "Nafta Super",
            "liters": 45.5,
            "station": "YPF",
            "notes": "Sin observaciones",
        }
    ]

    result_path = export_vouchers_to_excel(vouchers, output_path)

    assert result_path == output_path
    assert output_path.exists()

    workbook = load_workbook(output_path)
    sheet = workbook["Vouchers"]
    assert sheet.cell(1, 1).value == "Nro Vale"
    assert sheet.cell(2, 1).value == "VAL-202401-0001"


def test_create_database_backup_creates_copy(tmp_path: Path):
    source_path = tmp_path / "combustible.db"
    source_path.write_bytes(b"db-data")
    backup_dir = tmp_path / "backups"

    created_path = create_database_backup(source_path, backup_dir)

    assert created_path.exists()
    assert created_path.read_bytes() == b"db-data"
